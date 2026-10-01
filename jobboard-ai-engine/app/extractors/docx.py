import logging
from docx import Document
from io import BytesIO
from app.exceptions import ExtractionError

logger = logging.getLogger(__name__)


def _extract_hyperlink_targets(doc) -> list[str]:
    """Collect external hyperlink targets (mailto:, http(s)://) from the document
    body plus every section header/footer, in first-seen order. python-docx keeps
    these only as relationship targets — they never appear in paragraph text."""
    out: list[str] = []
    seen: set[str] = set()

    def _collect(part) -> None:
        rels = getattr(part, "rels", None)
        if not rels:
            return
        for rel in rels.values():
            try:
                if not rel.is_external:
                    continue
            except Exception:
                continue
            target = (rel.target_ref or "").strip()
            if target and target not in seen and (
                target.lower().startswith(("mailto:", "http://", "https://"))
            ):
                seen.add(target)
                out.append(target)

    _collect(doc.part)
    for section in doc.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            try:
                _collect(hf.part)
            except Exception:
                continue
    return out


def extract_text_from_docx(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from DOCX file, including tables. Returns (text, estimated_page_count).

    Raises ExtractionError on corrupt/empty. Page count estimated from char length (~3000 chars/page).
    """
    from app.config import settings

    try:
        doc = Document(BytesIO(file_bytes))
    except Exception as e:
        logger.error("DOCX extraction failed: %s", e)
        raise ExtractionError("Could not read DOCX file. It may be corrupted.") from e

    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    full_text = "\n".join(parts).strip()
    if not full_text:
        raise ExtractionError("No text found in DOCX file.")

    # Hyperlinks (email / LinkedIn / portfolio) live in relationship targets,
    # not the visible paragraph text — a "LinkedIn" label or an icon glyph links
    # out to the real URL. Surface external targets so contact extraction and the
    # LLM can read them, mirroring the PDF link-annotation path.
    links = _extract_hyperlink_targets(doc)
    new_links = []
    seen = set()
    for uri in links:
        if uri.lower().startswith("mailto:"):
            uri = uri[len("mailto:"):].split("?")[0].strip()  # drop mailto: and any ?subject=
        uri = uri.strip()
        if not uri or uri in seen or uri in full_text:
            continue
        seen.add(uri)
        new_links.append(uri)
    if new_links:
        full_text = (full_text + "\n\n[EMBEDDED LINKS]\n" + "\n".join(new_links)).strip()

    # Estimate page count (~3000 chars per page for typical resume formatting)
    page_count = max(1, len(full_text) // 3000 + (1 if len(full_text) % 3000 else 0))
    if page_count > settings.max_resume_pages:
        raise ExtractionError(
            f"Resume is approximately {page_count} pages. Maximum supported is {settings.max_resume_pages}."
        )

    return full_text, page_count

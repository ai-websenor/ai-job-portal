import logging
from docx import Document
from io import BytesIO
from app.exceptions import ExtractionError

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from DOCX file, including tables. Returns (text, estimated_page_count).

    Raises ExtractionError on corrupt/empty. Page count estimated from char length (~3000 chars/page).
    """
    from app.config import settings

    try:
        doc = Document(BytesIO(file_bytes))
    except Exception as e:
        logger.error("DOCX extraction failed: %s", e)
        raise ExtractionError("Could not read DOCX file. It may be corrupted.") from e

    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    full_text = "\n".join(parts).strip()
    if not full_text:
        raise ExtractionError("No text found in DOCX file.")

    # Estimate page count (~3000 chars per page for typical resume formatting)
    page_count = max(1, len(full_text) // 3000 + (1 if len(full_text) % 3000 else 0))
    if page_count > settings.max_resume_pages:
        raise ExtractionError(
            f"Resume is approximately {page_count} pages. Maximum supported is {settings.max_resume_pages}."
        )

    return full_text, page_count
